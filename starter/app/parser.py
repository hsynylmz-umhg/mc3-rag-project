"""Bounded, indexing-only extraction. No model or document library loads on import.

Each document is parsed in a disposable process: malformed native files cannot
hang the corpus walk. Text is evidence, never instructions or executable code.
"""
from __future__ import annotations

import csv
import io
import json
import math
import os
from pathlib import Path
import signal
import stat
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile

SUPPORTED = {".pdf", ".docx", ".xlsx", ".csv", ".tsv", ".txt", ".log", ".py", ".md", ".png", ".jpg", ".jpeg"}
MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_TEXT_CHARS = 4_000_000
MAX_PAGES = 500
MAX_ROWS = 40_000
MAX_COLUMNS = 256
MAX_IMAGE_PIXELS = 32_000_000


def _readable(path: Path) -> None:
    info = path.lstat()
    if stat.S_ISLNK(info.st_mode) or not stat.S_ISREG(info.st_mode):
        raise ValueError("symlink or non-regular file")
    # Root can open chmod(000) files: honor the corpus's unreadable-file intent.
    if info.st_mode & (stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH) == 0:
        raise PermissionError("no read permission bits")
    if info.st_size > MAX_FILE_BYTES:
        raise ValueError("file exceeds 64 MiB parsing limit")


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        if b"\x00" in data[:4096]:
            raise ValueError("binary content in text file")
        return data.decode("cp1252", errors="replace")


def _clean(text: str) -> str:
    return text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _blocks(text: str, location: str, size: int = 6000):
    """Split without losing a label at a chunk boundary; retain exact source text."""
    lines = _clean(text).splitlines()
    start = 0
    while start < len(lines):
        end, length = start, 0
        while end < len(lines) and (length < size or end == start):
            length += len(lines[end]) + 1
            end += 1
        piece = "\n".join(lines[start:end])
        if len(piece) > size * 2:
            for offset in range(0, len(piece), size - 300):
                yield f"{location}; lines {start + 1}-{end}; offset {offset}", piece[offset:offset + size]
        elif piece.strip():
            yield f"{location}; lines {start + 1}-{end}", piece
        if end == len(lines):
            break
        start = max(start + 1, end - 5)


def _check_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 10000 or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
            raise ValueError("office archive exceeds expansion limit")
        if any(m.file_size > 32 * 1024 * 1024 for m in members):
            raise ValueError("office archive member exceeds size limit")
        if any(m.flag_bits & 1 for m in members):
            raise ValueError("encrypted office archive")


def _table_row(values: list[str], headers: list[str]) -> str:
    return " | ".join(f"{headers[i] if i < len(headers) and headers[i] else 'column ' + str(i + 1)}: {value}"
                      for i, value in enumerate(values) if value)


def _parse_csv(path: Path):
    content = _decode(path.read_bytes())
    csv.field_size_limit(1_000_000)
    try:
        dialect = csv.Sniffer().sniff(content[:16000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = csv.reader(io.StringIO(content), dialect)
    headers: list[str] = []
    for n, row in enumerate(rows, 1):
        if n > MAX_ROWS:
            raise ValueError("CSV row limit reached")
        values = [_clean(v) for v in row[:MAX_COLUMNS]]
        if not any(values):
            continue
        if not headers:
            headers = values
            yield f"row {n}; headers", " | ".join(values)
        else:
            yield f"row {n}", _table_row(values, headers)


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if "%" in (cell.number_format or ""):
            return f"{value * 100:g}%"
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def _parse_xlsx(path: Path):
    import openpyxl

    _check_zip(path)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = None
    try:
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        for sheet in cached.worksheets:
            formula_sheet = formulas[sheet.title]
            # Ignore false worksheet dimensions found in some producer exports.
            sheet.reset_dimensions()
            formula_sheet.reset_dimensions()
            headers: list[str] = []
            titles: list[str] = []
            cached_rows = sheet.iter_rows(max_col=MAX_COLUMNS)
            formula_rows = formula_sheet.iter_rows(max_col=MAX_COLUMNS)
            for n, (row, formula_row) in enumerate(zip(cached_rows, formula_rows), 1):
                if n > MAX_ROWS:
                    raise ValueError("worksheet row limit reached")
                values = []
                for cell, formula in zip(row, formula_row):
                    value = _cell_text(cell)
                    if not value and formula.data_type == "f":
                        value = f"[formula with no cached result: {formula.value}]"
                    values.append(value)
                while values and not values[-1]:
                    values.pop()
                if not any(values):
                    continue
                prefix = f"Sheet: {sheet.title}\n" + ("\n".join(titles[-3:]) + "\n" if titles else "")
                if not headers and sum(bool(v) for v in values) >= 2:
                    headers = values
                    text = "Headers: " + " | ".join(values)
                elif not headers:
                    text = " | ".join(values)
                    titles.append(text)
                else:
                    text = _table_row(values, headers)
                yield f"sheet {sheet.title}; row {n}", prefix + text
    finally:
        cached.close()
        if formulas is not None:
            formulas.close()


def _parse_docx(path: Path):
    from defusedxml import ElementTree as ET

    _check_zip(path)
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def visible_text(node):
        result = []
        def visit(element):
            if element.tag in {ns + "del", ns + "moveFrom"}:
                return
            if element.tag == ns + "t":
                result.append(element.text or "")
            elif element.tag == ns + "tab":
                result.append("\t")
            elif element.tag in {ns + "br", ns + "cr"}:
                result.append("\n")
            for child in element:
                visit(child)
        visit(node)
        return _clean("".join(result))

    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n == "word/document.xml" or
                 (n.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")) and n.endswith(".xml"))]
        for name in sorted(names, key=lambda n: (n != "word/document.xml", n)):
            root = ET.fromstring(archive.read(name))
            paragraphs: list[str] = []
            heading = ""
            count = 0
            def walk(node):
                for child in node:
                    if child.tag in {ns + "p", ns + "tbl"}:
                        yield child
                    elif child.tag not in {ns + "del", ns + "moveFrom"}:
                        yield from walk(child)
            for element in walk(root):
                count += 1
                if element.tag == ns + "p":
                    text = visible_text(element)
                    if not text:
                        continue
                    paragraphs.append(text)
                    style = element.find(f"{ns}pPr/{ns}pStyle")
                    if style is not None and any(s in style.get(ns + "val", "").lower() for s in ("heading", "title")):
                        heading = text
                else:
                    if paragraphs:
                        yield from _blocks("\n".join(paragraphs), f"{name}; before block {count}")
                        paragraphs.clear()
                    headers = []
                    for row_number, row in enumerate(element.findall(ns + "tr"), 1):
                        values = [" ".join(visible_text(p) for p in cell.findall(ns + "p")) for cell in row.findall(ns + "tc")]
                        if not any(values):
                            continue
                        if not headers:
                            headers = values
                            text = "Headers: " + " | ".join(values)
                        else:
                            text = _table_row(values, headers)
                        yield f"{name}; table block {count}; row {row_number}", (heading + "\n" if heading else "") + text
            if paragraphs:
                yield from _blocks("\n".join(paragraphs), name)


def _ocr_layout(tsv: str, width: int) -> str:
    """Rebuild OCR rows from boxes, retaining diagram/table column alignment.

    Bounding boxes express only observed positions; no pin/signal relationships
    are invented. Words from separate Tesseract blocks may share a visual row.
    """
    words = []
    # Native Tesseract supplies a header; its WASM build can omit that header.
    header = "level page_num block_num par_num line_num word_num left top width height conf text".split()
    fields = None if tsv.startswith("level\t") else header
    for entry in csv.DictReader(io.StringIO(tsv), fieldnames=fields, delimiter="\t", quoting=csv.QUOTE_NONE):
        try:
            text = entry.get("text", "").strip()
            if entry.get("level") != "5" or not text or float(entry.get("conf", "-1")) < 0:
                continue
            left, top, height = int(entry["left"]), int(entry["top"]), int(entry["height"])
            words.append((top + height / 2, left, height, text))
        except (TypeError, ValueError, KeyError):
            continue
    rows: list[list] = []
    for word in sorted(words):
        if rows and abs(word[0] - statistics.median(w[0] for w in rows[-1])) <= min(word[2], statistics.median(w[2] for w in rows[-1])) * 0.6:
            rows[-1].append(word)
        else:
            rows.append([word])
    rendered = []
    # A common horizontal scale for EVERY row preserves vertical associations.
    for row in rows:
        line = ""
        for _, left, _, text in sorted(row, key=lambda word: word[1]):
            column = round(left / max(width, 1) * 150)
            line += " " * max(1 if line else 0, column - len(line)) + text
        # A non-whitespace gutter survives downstream strip()/chunk boundaries.
        rendered.append("| " + line.rstrip())
    return "\n".join(rendered)


def _remove_rules(image):
    """Remove long straight diagram/table rules without discarding text boxes.

    Tesseract otherwise treats the labels inside connector rectangles as rules
    and can omit every pin number. Thresholds exceed ordinary character strokes.
    """
    import numpy as np
    from PIL import Image, ImageFilter

    dark = np.asarray(image) < 160
    mask = np.zeros_like(dark)
    anchors = [[], []]
    for axis, (rows, target, minimum) in enumerate((
        (dark, mask, max(60, image.width // 25)),
        (dark.T, mask.T, max(60, image.height // 10)),
    )):
        for position, (row, marked) in enumerate(zip(rows, target)):
            edges = np.diff(np.pad(row.astype(np.int8), (1, 1)))
            starts, ends = np.flatnonzero(edges == 1), np.flatnonzero(edges == -1)
            for start, end in zip(starts, ends):
                if end - start >= minimum:
                    marked[start:end] = True
                if end - start >= (image.width * 0.5 if axis == 0 else image.height * 0.25):
                    anchors[axis].append(position)
    # Require a large frame/grid in BOTH directions. This avoids erasing tall
    # letter strokes in an ordinary cropped word or asset label.
    if any(len(values) < 2 for values in anchors):
        return image
    if max(anchors[0]) - min(anchors[0]) < image.height * 0.1 or max(anchors[1]) - min(anchors[1]) < image.width * 0.3:
        return image
    # Also remove antialias fringes along the detected rules.
    expanded = Image.fromarray(mask.astype(np.uint8) * 255).filter(ImageFilter.MaxFilter(3))
    result = image.copy()
    result.paste(255, mask=expanded)
    return result


def _ocr(image, *, psm: int = 6) -> str:
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(image).convert("RGB")
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise ValueError("image exceeds pixel limit")
    # Upscaling helps tiny printed labels; cap both memory and OCR latency.
    scale = min(2.0, 3000 / max(image.size))
    if scale != 1:
        image = image.resize((max(1, round(image.width * scale)), max(1, round(image.height * scale))), Image.Resampling.LANCZOS)
    image = ImageOps.autocontrast(ImageOps.grayscale(image))
    image = _remove_rules(image)
    with tempfile.TemporaryDirectory(prefix="mc3-ocr-") as folder:
        source = Path(folder) / "image.png"
        output_base = Path(folder) / "result"
        image.save(source)
        command = [os.environ.get("MC3_TESSERACT", "tesseract"), str(source), str(output_base), "-l",
                   os.environ.get("MC3_OCR_LANG", "eng"), "--psm", str(psm), "-c", "preserve_interword_spaces=1", "txt", "tsv"]
        result = subprocess.run(command, capture_output=True, timeout=12, check=False,
                                env={**os.environ, "OMP_THREAD_LIMIT": "1"})
        if result.returncode != 0:
            raise RuntimeError("tesseract failed: " + result.stderr.decode("utf-8", "replace")[-400:])
        layout = _ocr_layout(output_base.with_suffix(".tsv").read_text(encoding="utf-8"), image.width)
        return layout or _clean(output_base.with_suffix(".txt").read_text(encoding="utf-8"))


def _parse_image(path: Path):
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    with Image.open(path) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ValueError("image exceeds pixel limit")
        text = _ocr(image, psm=6)
        if not text:
            text = _ocr(image, psm=11)
        if not text:
            raise ValueError("OCR found no text")
        yield from _blocks(text, "image OCR")


def _parse_pdf(path: Path):
    import fitz
    from PIL import Image

    with fitz.open(path) as document:
        # Even an empty user password is not an invitation to extract protected data.
        if document.is_encrypted or document.needs_pass or document.xref_get_key(-1, "Encrypt")[0] != "null":
            raise ValueError("encrypted PDF skipped")
        for n, page in enumerate(document, 1):
            if n > MAX_PAGES:
                raise ValueError("PDF page limit reached")
            text = page.get_text("text", sort=True)
            sparse = len("".join(text.split())) < 25
            try:
                page_area = max(1, page.rect.width * page.rect.height)
                has_diagram = any(fitz.Rect(info["bbox"]).get_area() > page_area * 0.1 for info in page.get_image_info())
                if sparse or has_diagram:
                    zoom = min(2.2, (MAX_IMAGE_PIXELS / page_area) ** 0.5)
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), colorspace=fitz.csRGB, alpha=False)
                    scanned = _ocr(Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples))
                    if scanned:
                        text = scanned if sparse else text + "\n[Page image OCR; columns retain visual alignment]\n" + scanned
            except Exception:
                # Optional image discovery/rendering/OCR must not lose native text.
                if sparse:
                    raise
            yield from _blocks(text, f"page {n}")


def _extract(path: Path):
    _readable(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        yield from _parse_pdf(path)
    elif suffix == ".docx":
        yield from _parse_docx(path)
    elif suffix == ".xlsx":
        yield from _parse_xlsx(path)
    elif suffix in {".csv", ".tsv"}:
        yield from _parse_csv(path)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        yield from _parse_image(path)
    elif suffix in SUPPORTED:
        yield from _blocks(_decode(path.read_bytes()), "text")
    else:
        raise ValueError("unsupported file format")


def _worker(path: Path, relative: str) -> dict:
    # Resource limits are Unix-only; the parent timeout also works on Windows.
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
        resource.setrlimit(resource.RLIMIT_CPU, (25, 26))
        resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_EXPANDED_BYTES, MAX_EXPANDED_BYTES))
    except (ImportError, ValueError, OSError):
        pass
    units, errors, total = [], [], 0
    try:
        for location, text in _extract(path):
            text = _clean(text)
            if not text:
                continue
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise ValueError("extracted text limit reached")
            units.append({"path": relative, "location": location, "text": text})
    except Exception as exc:
        errors.append({"path": relative, "error": f"{type(exc).__name__}: {exc}"[:600]})
    return {"units": units, "errors": errors}


def _kill_worker(process: subprocess.Popen) -> None:
    try:
        if os.name == "posix":
            os.killpg(process.pid, signal.SIGKILL)
        else:
            # Includes a possible OCR descendant; no shell and no user path interpolation.
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=3, check=False)
    except (OSError, subprocess.SubprocessError):
        process.kill()
    process.communicate(timeout=3)


def parse_corpus(corpus: Path, *, file_timeout: float = 25.0,
                 total_timeout: float = 210.0) -> tuple[list[dict], list[dict]]:
    """Return source units plus diagnostic errors, with a deadline for the whole walk.

    All paths are corpus-relative POSIX paths. Nothing executes source documents,
    opens encrypted files with passwords, follows links, or modifies the corpus.
    """
    corpus = Path(corpus).resolve()
    units: list[dict] = []
    errors: list[dict] = []
    deadline = time.monotonic() + total_timeout
    if not corpus.is_dir():
        return units, [{"path": ".", "error": "corpus root is not a directory"}]

    def walk_error(exc):
        errors.append({"path": str(getattr(exc, "filename", ".")), "error": str(exc)[:600]})

    for folder, directories, filenames in os.walk(corpus, followlinks=False, onerror=walk_error):
        directories[:] = sorted(d for d in directories if not (Path(folder) / d).is_symlink())
        for name in sorted(filenames):
            path = Path(folder) / name
            relative = path.relative_to(corpus).as_posix()
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                errors.append({"path": ".", "error": "corpus parsing deadline reached; remaining files skipped"})
                return units, errors
            try:
                _readable(path)
                if path.suffix.lower() not in SUPPORTED:
                    raise ValueError("unsupported file format")
                process = subprocess.Popen([sys.executable, str(Path(__file__).resolve()), "--parse-file", str(path), relative],
                                           stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                           start_new_session=(os.name == "posix"),
                                           env={**os.environ, "PYTHONIOENCODING": "utf-8", "OMP_NUM_THREADS": "1",
                                                "OPENBLAS_NUM_THREADS": "1"})
                try:
                    output, error_output = process.communicate(timeout=min(file_timeout, remaining))
                except subprocess.TimeoutExpired:
                    _kill_worker(process)
                    raise TimeoutError("document parsing time limit reached") from None
                if process.returncode:
                    raise RuntimeError(f"parser exited {process.returncode}: {error_output.decode('utf-8', 'replace')[-400:]}")
                payload = json.loads(output)
                units.extend(payload["units"])
                errors.extend(payload["errors"])
            except Exception as exc:
                errors.append({"path": relative, "error": f"{type(exc).__name__}: {exc}"[:600]})
    return units, errors


if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--parse-file":
        print(json.dumps(_worker(Path(sys.argv[2]), sys.argv[3]), ensure_ascii=True))
    else:
        raise SystemExit("This module is used only by the indexing process.")
